# -*- coding: utf-8 -*-
"""
Gera a planilha de controle de gastos da viagem ao Japao.

Estrutura do arquivo gerado:
  Painel   -> resumo visual, cotacoes, orcamento por categoria
  Gastos   -> lancamento dos gastos (a folha que se preenche no dia a dia)
  Por Dia  -> quanto foi gasto em cada dia da viagem
  Listas   -> listas que alimentam os menus suspensos
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from datetime import date

OUT = "Controle_de_Gastos_Japao.xlsx"

# ---------------------------------------------------------------- paleta
INK        = "1C2541"   # azul marinho (cabecalhos)
ACCENT     = "C1121F"   # vermelho (acento japones)
TEAL       = "1B7F5A"   # verde
AMBER      = "B26B00"   # ambar
BLUE       = "2A5DB0"   # azul
GREY       = "6B7280"   # texto secundario
CARD_BG    = "F3F5F9"
BAND       = "FAFBFD"
LINE       = "DCE0E8"
INPUT_BG   = "FFF6D6"   # amarelo = celula para preencher
INPUT_FG   = "0000FF"   # azul = valor digitado
CALC_BG    = "EEF1F6"   # cinza = calculado automaticamente
OK_BG      = "DDF0E5"
WARN_BG    = "FBEBD2"
INFO_BG    = "E1E9F7"
MUTE_BG    = "ECEDF0"

FONT = "Arial"

# ---------------------------------------------------------------- formatos
F_BRL   = 'R$ #,##0.00;-R$ #,##0.00;"-"'
F_NUM   = '#,##0.00;-#,##0.00;"-"'
F_INT   = '#,##0;-#,##0;"-"'
F_PCT   = '0.0%'
F_DATE  = 'DD/MM/YYYY'
F_RATE  = '#,##0.0000'

thin = Side(style="thin", color=LINE)


def f(size=10, bold=False, color="1C2541", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def left(wrap=False, indent=0):
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap, indent=indent)


def right():
    return Alignment(horizontal="right", vertical="center")


def outline(ws, c1, r1, c2, r2, color=LINE, style="thin"):
    """Desenha uma moldura ao redor de um bloco de celulas."""
    s = Side(style=style, color=color)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            b = cell.border
            cell.border = Border(
                left=s if c == c1 else b.left,
                right=s if c == c2 else b.right,
                top=s if r == r1 else b.top,
                bottom=s if r == r2 else b.bottom,
            )


def paint(ws, c1, r1, c2, r2, patt):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).fill = patt


def top_accent(ws, c1, r1, c2, color):
    """Linha colorida no topo de um cartao."""
    s = Side(style="medium", color=color)
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=r1, column=c)
        b = cell.border
        cell.border = Border(left=b.left, right=b.right, top=s, bottom=b.bottom)


# ---------------------------------------------------------------- dados base
CATEGORIAS = [
    "Passagens Aereas",
    "Hospedagem",
    "Transporte entre cidades (JR/Shinkansen)",
    "Transporte local (metro/onibus/taxi)",
    "Alimentacao",
    "Passeios e Ingressos",
    "Compras e Souvenirs",
    "Seguro Viagem",
    "Documentos e Vistos",
    "Chip / Internet",
    "Saude e Farmacia",
    "Bagagem e Equipamentos",
    "Outros",
]
# acentuacao correta (mantida separada para facilitar leitura do codigo acima)
CATEGORIAS = [
    "Passagens Aéreas",
    "Hospedagem",
    "Transporte entre cidades (JR/Shinkansen)",
    "Transporte local (metrô/ônibus/táxi)",
    "Alimentação",
    "Passeios e Ingressos",
    "Compras e Souvenirs",
    "Seguro Viagem",
    "Documentos e Vistos",
    "Chip / Internet",
    "Saúde e Farmácia",
    "Bagagem e Equipamentos",
    "Outros",
]

ORCAMENTO = [12000, 9000, 3500, 1200, 6000, 2500, 3000, 700, 300, 400, 300, 800, 1000]

PAGAMENTOS = [
    "Dinheiro (iene)",
    "Dinheiro (real)",
    "Cartão de crédito",
    "Cartão de débito",
    "Cartão internacional (Wise/Nomad)",
    "IC Card (Suica/Pasmo)",
    "PIX",
    "Transferência / Boleto",
    "Voucher pré-pago",
]

STATUS = ["Pago", "Pendente", "Parcelado", "Reembolsado"]

PESSOAS = ["Viajante 1", "Viajante 2", "Dividido", "Grupo"]

CIDADES = [
    "Tóquio", "Quioto", "Osaka", "Nara", "Hiroshima", "Hakone", "Nikko",
    "Kanazawa", "Takayama", "Sapporo", "Fukuoka", "Monte Fuji",
    "Brasil (pré-viagem)", "Outro",
]

MOEDAS = [
    ("JPY", 0.0380, "Iene japonês"),
    ("BRL", 1.0000, "Real brasileiro"),
    ("USD", 5.4000, "Dólar americano"),
    ("EUR", 5.9000, "Euro"),
]

FIRST = 4          # primeira linha de dados em "Gastos"
LAST = 400         # ultima linha preparada em "Gastos"

wb = Workbook()

# =================================================================== PAINEL
pa = wb.active
pa.title = "Painel"
pa.sheet_view.showGridLines = False
pa.sheet_properties.tabColor = INK

widths = {"A": 2.5, "B": 30, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 16, "I": 16, "J": 2.5}
for col, w in widths.items():
    pa.column_dimensions[col].width = w

# --- faixa / titulo
pa.row_dimensions[1].height = 6
pa.row_dimensions[2].height = 7
paint(pa, 2, 2, 9, 2, fill(ACCENT))

pa.row_dimensions[3].height = 38
pa.row_dimensions[4].height = 22
pa.merge_cells("B3:I3")
pa.merge_cells("B4:I4")
paint(pa, 2, 3, 9, 4, fill(INK))
pa["B3"] = "CONTROLE DE GASTOS  •  VIAGEM AO JAPÃO"
pa["B3"].font = f(20, True, "FFFFFF")
pa["B3"].alignment = left(indent=1)
pa["B4"] = "Todos os gastos da viagem em um lugar só — planejado, pago e em aberto"
pa["B4"].font = f(10, False, "C9CEDD")
pa["B4"].alignment = left(indent=1)
pa.row_dimensions[5].height = 12

# --- blocos: dados da viagem  |  cambio
def section(ws, cell_range, texto, cor=INK):
    ws.merge_cells(cell_range)
    first = cell_range.split(":")[0]
    ws[first] = texto
    ws[first].font = f(10, True, "FFFFFF")
    ws[first].alignment = left(indent=1)
    c1 = ws[first].column
    r1 = ws[first].row
    c2 = ws[cell_range.split(":")[1]].column
    paint(ws, c1, r1, c2, r1, fill(cor))
    ws.row_dimensions[r1].height = 20


section(pa, "B6:C6", "DADOS DA VIAGEM")
section(pa, "E6:G6", "CÂMBIO — quanto vale 1 unidade em R$", ACCENT)

info = [
    ("Destino", "Japão", "text"),
    ("Data de ida", date(2026, 10, 3), "date"),
    ("Data de volta", date(2026, 10, 17), "date"),
    ("Nº de viajantes", 2, "int"),
]
r = 7
for label, val, kind in info:
    pa.cell(row=r, column=2, value=label).font = f(10, False, GREY)
    pa.cell(row=r, column=2).alignment = left(indent=1)
    c = pa.cell(row=r, column=3, value=val)
    c.font = f(10, True, INPUT_FG)
    c.fill = fill(INPUT_BG)
    c.alignment = center()
    c.number_format = {"date": F_DATE, "int": F_INT, "text": "General"}[kind]
    pa.row_dimensions[r].height = 19
    r += 1

pa.cell(row=r, column=2, value="Duração (dias)").font = f(10, False, GREY)
pa.cell(row=r, column=2).alignment = left(indent=1)
c = pa.cell(row=r, column=3, value="=IFERROR($C$9-$C$8+1,\"\")")
c.font = f(10, True, INK); c.fill = fill(CALC_BG); c.alignment = center(); c.number_format = F_INT
pa.row_dimensions[r].height = 19
r += 1
pa.cell(row=r, column=2, value="Orçamento total (R$)").font = f(10, False, GREY)
pa.cell(row=r, column=2).alignment = left(indent=1)
c = pa.cell(row=r, column=3, value="=SUM($C$25:$C$37)")
c.font = f(10, True, INK); c.fill = fill(CALC_BG); c.alignment = center(); c.number_format = F_BRL
pa.row_dimensions[r].height = 19
outline(pa, 2, 7, 3, 12)

# cabecalho da tabela de cambio
for col, txt in ((5, "Moeda"), (6, "Cotação (R$)"), (7, "Descrição")):
    cell = pa.cell(row=7, column=col, value=txt)
    cell.font = f(9, True, GREY)
    cell.alignment = center()
    cell.fill = fill(BAND)
pa.row_dimensions[7].height = 19

for i, (cod, taxa, desc) in enumerate(MOEDAS):
    rr = 8 + i
    pa.cell(row=rr, column=5, value=cod).font = f(10, True, INK)
    pa.cell(row=rr, column=5).alignment = center()
    cc = pa.cell(row=rr, column=6, value=taxa)
    cc.font = f(10, True, INPUT_FG); cc.fill = fill(INPUT_BG)
    cc.alignment = center(); cc.number_format = F_RATE
    pa.cell(row=rr, column=7, value=desc).font = f(10, False, GREY)
    pa.cell(row=rr, column=7).alignment = left(indent=1)
    pa.row_dimensions[rr].height = 19

pa.merge_cells("E12:G12")
pa["E12"] = "Atualize a cotação antes e durante a viagem."
pa["E12"].font = f(9, False, GREY, italic=True)
pa["E12"].alignment = left(indent=1)
outline(pa, 5, 7, 7, 11)

pa.row_dimensions[13].height = 12

# --- KPIs
section(pa, "B14:I14", "RESUMO")

KPIS = [
    # (linha_label, coluna_inicial, titulo, formula, formato, cor_acento, fundo)
    (16, 2, "TOTAL GASTO",
     f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$G${FIRST}:$G${LAST},"<>Reembolsado")', F_BRL, INK, CARD_BG),
    (16, 4, "JÁ PAGO",
     f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$G${FIRST}:$G${LAST},"Pago")', F_BRL, TEAL, OK_BG),
    (16, 6, "EM ABERTO",
     f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$G${FIRST}:$G${LAST},"Pendente")'
     f'+SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$G${FIRST}:$G${LAST},"Parcelado")', F_BRL, AMBER, WARN_BG),
    (16, 8, "SALDO DO ORÇAMENTO",
     f'=$C$12-SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$G${FIRST}:$G${LAST},"<>Reembolsado")',
     F_BRL, BLUE, INFO_BG),
    (19, 2, "GASTO POR PESSOA", '=IFERROR($B$17/$C$10,0)', F_BRL, INK, CARD_BG),
    (19, 4, "GASTO POR DIA DE VIAGEM", '=IFERROR($B$17/$C$11,0)', F_BRL, INK, CARD_BG),
    (19, 6, "GASTOS PRÉ-VIAGEM",
     f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$A${FIRST}:$A${LAST},"<"&$C$8)', F_BRL, GREY, MUTE_BG),
    (19, 8, "GASTOS DURANTE A VIAGEM",
     f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$A${FIRST}:$A${LAST},">="&$C$8)', F_BRL, ACCENT, CARD_BG),
]

for rl, c1, titulo, formula, fmt, cor, bg in KPIS:
    c2 = c1 + 1
    pa.merge_cells(start_row=rl, start_column=c1, end_row=rl, end_column=c2)
    pa.merge_cells(start_row=rl + 1, start_column=c1, end_row=rl + 1, end_column=c2)
    paint(pa, c1, rl, c2, rl + 1, fill(bg))
    t = pa.cell(row=rl, column=c1, value=titulo)
    t.font = f(8, True, GREY)
    t.alignment = center(wrap=True)
    v = pa.cell(row=rl + 1, column=c1, value=formula)
    v.font = f(15, True, cor)
    v.alignment = center()
    v.number_format = fmt
    outline(pa, c1, rl, c2, rl + 1)
    top_accent(pa, c1, rl, c2, cor)
    pa.row_dimensions[rl].height = 20
    pa.row_dimensions[rl + 1].height = 30

pa.row_dimensions[18].height = 8
pa.row_dimensions[21].height = 14

# --- tabela por categoria
section(pa, "B23:I23", "GASTOS POR CATEGORIA")

heads = ["Categoria", "Orçamento (R$)", "Gasto (R$)", "Já pago", "Em aberto",
         "% do total", "Saldo do orçamento", "Lançamentos"]
for i, h in enumerate(heads):
    cell = pa.cell(row=24, column=2 + i, value=h)
    cell.font = f(9, True, "FFFFFF")
    cell.fill = fill("3B4A6B")
    cell.alignment = center(wrap=True)
pa.row_dimensions[24].height = 28

for i, cat in enumerate(CATEGORIAS):
    rr = 25 + i
    pa.row_dimensions[rr].height = 18
    band = fill(BAND) if i % 2 == 0 else fill("FFFFFF")

    a = pa.cell(row=rr, column=2, value=f"=IF(Listas!$B${5+i}=\"\",\"\",Listas!$B${5+i})")
    a.font = f(10, False, INK); a.alignment = left(indent=1); a.fill = band

    b = pa.cell(row=rr, column=3, value=ORCAMENTO[i])
    b.font = f(10, False, INPUT_FG); b.fill = fill(INPUT_BG)
    b.alignment = right(); b.number_format = F_BRL

    for col, formula in (
        (4, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$B${FIRST}:$B${LAST},$B{rr},'
            f'Gastos!$G${FIRST}:$G${LAST},"<>Reembolsado")'),
        (5, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$B${FIRST}:$B${LAST},$B{rr},'
            f'Gastos!$G${FIRST}:$G${LAST},"Pago")'),
        (6, f'=$D{rr}-$E{rr}'),
        (7, f'=IFERROR($D{rr}/$D$38,0)'),
        (8, f'=$C{rr}-$D{rr}'),
        (9, f'=COUNTIFS(Gastos!$B${FIRST}:$B${LAST},$B{rr})'),
    ):
        cell = pa.cell(row=rr, column=col, value=formula)
        cell.font = f(10, False, INK)
        cell.alignment = right()
        cell.fill = band
        cell.number_format = F_PCT if col == 7 else (F_INT if col == 9 else F_BRL)

# total
tr = 38
pa.row_dimensions[tr].height = 22
pa.cell(row=tr, column=2, value="TOTAL").font = f(10, True, "FFFFFF")
pa.cell(row=tr, column=2).alignment = left(indent=1)
for col in range(3, 10):
    L = get_column_letter(col)
    cell = pa.cell(row=tr, column=col, value=f"=SUM({L}25:{L}37)")
    cell.font = f(10, True, "FFFFFF")
    cell.alignment = right()
    cell.number_format = F_INT if col == 9 else F_BRL
pa.cell(row=tr, column=7, value="=IFERROR($D$38/$D$38,0)").number_format = F_PCT
pa.cell(row=tr, column=7).font = f(10, True, "FFFFFF")
pa.cell(row=tr, column=7).alignment = right()
paint(pa, 2, tr, 9, tr, fill(INK))
outline(pa, 2, 24, 9, tr)

pa.conditional_formatting.add(
    f"D25:D37",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=ACCENT, showValue=True),
)
pa.conditional_formatting.add(
    f"H25:H37",
    CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, size=10, bold=True, color=ACCENT)),
)

pa.row_dimensions[39].height = 12

# --- por forma de pagamento  |  por pessoa
section(pa, "B40:E40", "POR FORMA DE PAGAMENTO", "3B4A6B")
section(pa, "G40:I40", "POR PESSOA", "3B4A6B")

for i, h in enumerate(["Forma de pagamento", "Total (R$)", "Já pago", "Em aberto"]):
    cell = pa.cell(row=41, column=2 + i, value=h)
    cell.font = f(9, True, GREY); cell.fill = fill(BAND); cell.alignment = center(wrap=True)
for i, h in enumerate(["Quem pagou", "Total (R$)", "Em aberto"]):
    cell = pa.cell(row=41, column=7 + i, value=h)
    cell.font = f(9, True, GREY); cell.fill = fill(BAND); cell.alignment = center(wrap=True)
pa.row_dimensions[41].height = 24

for i in range(len(PAGAMENTOS)):
    rr = 42 + i
    pa.row_dimensions[rr].height = 18
    a = pa.cell(row=rr, column=2, value=f"=IF(Listas!$D${5+i}=\"\",\"\",Listas!$D${5+i})")
    a.font = f(10, False, INK); a.alignment = left(indent=1)
    for col, formula in (
        (3, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$F${FIRST}:$F${LAST},$B{rr})'),
        (4, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$F${FIRST}:$F${LAST},$B{rr},'
            f'Gastos!$G${FIRST}:$G${LAST},"Pago")'),
        (5, f'=$C{rr}-$D{rr}'),
    ):
        cell = pa.cell(row=rr, column=col, value=formula)
        cell.font = f(10, False, INK); cell.alignment = right(); cell.number_format = F_BRL
outline(pa, 2, 41, 5, 41 + len(PAGAMENTOS))

for i in range(len(PESSOAS)):
    rr = 42 + i
    a = pa.cell(row=rr, column=7, value=f"=IF(Listas!$H${5+i}=\"\",\"\",Listas!$H${5+i})")
    a.font = f(10, False, INK); a.alignment = left(indent=1)
    for col, formula in (
        (8, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$E${FIRST}:$E${LAST},$G{rr})'),
        (9, f'=SUMIFS(Gastos!$J${FIRST}:$J${LAST},Gastos!$E${FIRST}:$E${LAST},$G{rr},'
            f'Gastos!$G${FIRST}:$G${LAST},"Pendente")'),
    ):
        cell = pa.cell(row=rr, column=col, value=formula)
        cell.font = f(10, False, INK); cell.alignment = right(); cell.number_format = F_BRL
outline(pa, 7, 41, 9, 41 + len(PESSOAS))

# --- como usar
guia_r = 42 + len(PAGAMENTOS) + 1
section(pa, f"B{guia_r}:I{guia_r}", "COMO USAR ESTA PLANILHA", ACCENT)
linhas = [
    "1.  Preencha os dados da viagem e as cotações do dia aqui no Painel (células amarelas).",
    "2.  Ajuste o orçamento de cada categoria na coluna \"Orçamento (R$)\" da tabela acima.",
    "3.  Lance cada gasto na aba GASTOS, uma linha por gasto. As colunas com menu suspenso já vêm prontas.",
    "4.  Digite o valor na moeda em que você pagou e escolha a moeda — a conversão para real é automática.",
    "5.  Marque o STATUS: Pago, Pendente, Parcelado ou Reembolsado. Este Painel se atualiza sozinho.",
    "6.  A aba POR DIA mostra quanto saiu em cada dia da viagem; a aba LISTAS é onde você edita as opções dos menus.",
    "•  Amarelo = você preenche.   Cinza = calculado automaticamente, não precisa mexer.",
]
for i, t in enumerate(linhas):
    rr = guia_r + 1 + i
    pa.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=9)
    cell = pa.cell(row=rr, column=2, value=t)
    cell.font = f(9, i == len(linhas) - 1, GREY)
    cell.alignment = left(indent=1)
    pa.row_dimensions[rr].height = 17
    paint(pa, 2, rr, 9, rr, fill(BAND))
outline(pa, 2, guia_r + 1, 9, guia_r + len(linhas))

# =================================================================== GASTOS
gs = wb.create_sheet("Gastos")
gs.sheet_view.showGridLines = False
gs.sheet_properties.tabColor = ACCENT

COLS = [
    ("Data", 12, F_DATE, "input"),
    ("Categoria", 30, "General", "input"),
    ("Descrição", 34, "General", "input"),
    ("Cidade / Local", 18, "General", "input"),
    ("Quem pagou", 15, "General", "input"),
    ("Forma de pagamento", 24, "General", "input"),
    ("Status", 13, "General", "input"),
    ("Moeda", 9, "General", "input"),
    ("Valor pago", 14, F_NUM, "input"),
    ("Valor em R$", 15, F_BRL, "calc"),
    ("Observações", 34, "General", "input"),
]
for i, (h, w, _, _) in enumerate(COLS):
    gs.column_dimensions[get_column_letter(i + 1)].width = w

gs.row_dimensions[1].height = 32
gs.merge_cells("A1:K1")
gs["A1"] = "LANÇAMENTO DE GASTOS  •  VIAGEM AO JAPÃO"
gs["A1"].font = f(15, True, "FFFFFF")
gs["A1"].alignment = left(indent=1)
paint(gs, 1, 1, 11, 1, fill(INK))

gs.row_dimensions[2].height = 20
gs.merge_cells("A2:K2")
gs["A2"] = ("Uma linha por gasto.  Colunas com seta abrem menu suspenso.  "
            "\"Valor em R$\" é calculado sozinho pela cotação do Painel — não digite nada nela.")
gs["A2"].font = f(9, False, GREY, italic=True)
gs["A2"].alignment = left(indent=1)
paint(gs, 1, 2, 11, 2, fill(BAND))

gs.row_dimensions[3].height = 26
for i, (h, _, _, kind) in enumerate(COLS):
    cell = gs.cell(row=3, column=i + 1, value=h)
    cell.font = f(9, True, "FFFFFF")
    cell.fill = fill(ACCENT if kind == "calc" else "3B4A6B")
    cell.alignment = center(wrap=True)

exemplo = [
    date(2026, 9, 2), "Hospedagem", "Hotel em Shinjuku — 4 noites", "Tóquio",
    "Viajante 1", "Cartão de crédito", "Pago", "JPY", 86000, None,
    "Linha de exemplo — apague quando começar a preencher",
]

for rr in range(FIRST, LAST + 1):
    gs.row_dimensions[rr].height = 18
    band = fill("FFFFFF") if (rr - FIRST) % 2 == 0 else fill(BAND)
    for i, (_, _, fmt, kind) in enumerate(COLS):
        col = i + 1
        cell = gs.cell(row=rr, column=col)
        cell.font = f(10, False, INK)
        cell.number_format = fmt
        cell.fill = fill(CALC_BG) if kind == "calc" else band
        cell.alignment = right() if col in (9, 10) else (center() if col in (1, 7, 8) else left(indent=1))
        cell.border = Border(bottom=thin)
    # conversao automatica para real
    gs.cell(row=rr, column=10).value = (
        f'=IF($I{rr}="","",IFERROR($I{rr}*INDEX(TAXAS,MATCH($H{rr},MOEDAS,0)),""))'
    )
    if rr == FIRST:
        for i, v in enumerate(exemplo):
            if v is not None:
                gs.cell(row=rr, column=i + 1).value = v
        gs.cell(row=rr, column=11).font = f(9, False, GREY, italic=True)

gs.freeze_panes = "A4"
gs.auto_filter.ref = f"A3:K{LAST}"

# menus suspensos
dvs = [
    ("CATEGORIAS", f"B{FIRST}:B{LAST}", "Escolha uma categoria da lista (aba Listas)."),
    ("CIDADES", f"D{FIRST}:D{LAST}", "Escolha a cidade/local (aba Listas)."),
    ("PESSOAS", f"E{FIRST}:E{LAST}", "Quem pagou este gasto?"),
    ("PAGAMENTOS", f"F{FIRST}:F{LAST}", "Como foi pago?"),
    ("STATUS", f"G{FIRST}:G{LAST}", "Pago, Pendente, Parcelado ou Reembolsado."),
    ("MOEDAS", f"H{FIRST}:H{LAST}", "Moeda em que o valor foi pago."),
]
for nome, ref, msg in dvs:
    dv = DataValidation(type="list", formula1=f"={nome}", allow_blank=True, showDropDown=False)
    dv.error = "Escolha uma das opções da lista."
    dv.errorTitle = "Valor inválido"
    dv.prompt = msg
    dv.promptTitle = nome.capitalize()
    gs.add_data_validation(dv)
    dv.add(ref)

# cores por status
gs.conditional_formatting.add(
    f"G{FIRST}:G{LAST}",
    CellIsRule(operator="equal", formula=['"Pago"'], fill=fill(OK_BG),
               font=Font(name=FONT, size=10, bold=True, color=TEAL)))
gs.conditional_formatting.add(
    f"G{FIRST}:G{LAST}",
    CellIsRule(operator="equal", formula=['"Pendente"'], fill=fill(WARN_BG),
               font=Font(name=FONT, size=10, bold=True, color=AMBER)))
gs.conditional_formatting.add(
    f"G{FIRST}:G{LAST}",
    CellIsRule(operator="equal", formula=['"Parcelado"'], fill=fill(INFO_BG),
               font=Font(name=FONT, size=10, bold=True, color=BLUE)))
gs.conditional_formatting.add(
    f"G{FIRST}:G{LAST}",
    CellIsRule(operator="equal", formula=['"Reembolsado"'], fill=fill(MUTE_BG),
               font=Font(name=FONT, size=10, bold=True, color=GREY)))

# =================================================================== POR DIA
pd_ = wb.create_sheet("Por Dia")
pd_.sheet_view.showGridLines = False
pd_.sheet_properties.tabColor = TEAL

DIA_COLS = [("Data", 13), ("Total do dia (R$)", 16), ("Acumulado (R$)", 16),
            ("Alimentação", 15), ("Transporte local", 16), ("Passeios e Ingressos", 18),
            ("Compras", 19), ("Demais categorias", 17), ("Lançamentos", 12)]
for i, (h, w) in enumerate(DIA_COLS):
    pd_.column_dimensions[get_column_letter(i + 1)].width = w

pd_.row_dimensions[1].height = 32
pd_.merge_cells("A1:I1")
pd_["A1"] = "GASTOS DIA A DIA"
pd_["A1"].font = f(15, True, "FFFFFF")
pd_["A1"].alignment = left(indent=1)
paint(pd_, 1, 1, 9, 1, fill(INK))

pd_.row_dimensions[2].height = 20
pd_.merge_cells("A2:I2")
pd_["A2"] = "As datas vêm das datas de ida e volta do Painel. Tudo aqui é calculado automaticamente."
pd_["A2"].font = f(9, False, GREY, italic=True)
pd_["A2"].alignment = left(indent=1)
paint(pd_, 1, 2, 9, 2, fill(BAND))

pd_.row_dimensions[3].height = 28
for i, (h, _) in enumerate(DIA_COLS):
    cell = pd_.cell(row=3, column=i + 1, value=h)
    cell.font = f(9, True, "FFFFFF")
    cell.fill = fill("3B4A6B")
    cell.alignment = center(wrap=True)
# cabecalhos de categoria ligados a aba Listas (servem de critério do SUMIFS)
pd_["D3"] = "=Listas!$B$9"
pd_["E3"] = "=Listas!$B$8"
pd_["F3"] = "=Listas!$B$10"
pd_["G3"] = "=Listas!$B$11"

DIA_FIRST, DIA_LAST = 4, 43
for rr in range(DIA_FIRST, DIA_LAST + 1):
    pd_.row_dimensions[rr].height = 18
    band = fill("FFFFFF") if (rr - DIA_FIRST) % 2 == 0 else fill(BAND)

    if rr == DIA_FIRST:
        pd_.cell(row=rr, column=1, value='=IF(Painel!$C$8="","",Painel!$C$8)')
    else:
        pd_.cell(row=rr, column=1,
                 value=f'=IF($A{rr-1}="","",IF($A{rr-1}+1>Painel!$C$9,"",$A{rr-1}+1))')

    pd_.cell(row=rr, column=2,
             value=f'=IF($A{rr}="","",SUMIFS(Gastos!$J${FIRST}:$J${LAST},'
                   f'Gastos!$A${FIRST}:$A${LAST},$A{rr}))')
    pd_.cell(row=rr, column=3, value=f'=IF($A{rr}="","",SUM($B${DIA_FIRST}:$B{rr}))')
    for col in (4, 5, 6, 7):
        L = get_column_letter(col)
        pd_.cell(row=rr, column=col,
                 value=f'=IF($A{rr}="","",SUMIFS(Gastos!$J${FIRST}:$J${LAST},'
                       f'Gastos!$A${FIRST}:$A${LAST},$A{rr},'
                       f'Gastos!$B${FIRST}:$B${LAST},{L}$3))')
    pd_.cell(row=rr, column=8,
             value=f'=IF($A{rr}="","",$B{rr}-SUM($D{rr}:$G{rr}))')
    pd_.cell(row=rr, column=9,
             value=f'=IF($A{rr}="","",COUNTIFS(Gastos!$A${FIRST}:$A${LAST},$A{rr}))')

    for col in range(1, 10):
        cell = pd_.cell(row=rr, column=col)
        cell.font = f(10, col == 2, INK)
        cell.fill = band
        cell.border = Border(bottom=thin)
        cell.alignment = center() if col in (1, 9) else right()
        cell.number_format = F_DATE if col == 1 else (F_INT if col == 9 else F_BRL)

tot = DIA_LAST + 1
pd_.row_dimensions[tot].height = 22
pd_.cell(row=tot, column=1, value="TOTAL").font = f(10, True, "FFFFFF")
pd_.cell(row=tot, column=1).alignment = center()
for col in range(2, 10):
    L = get_column_letter(col)
    if col == 3:
        continue
    cell = pd_.cell(row=tot, column=col, value=f"=SUM({L}{DIA_FIRST}:{L}{DIA_LAST})")
    cell.font = f(10, True, "FFFFFF")
    cell.alignment = right()
    cell.number_format = F_INT if col == 9 else F_BRL
paint(pd_, 1, tot, 9, tot, fill(INK))
outline(pd_, 1, 3, 9, tot)

pd_.conditional_formatting.add(
    f"B{DIA_FIRST}:B{DIA_LAST}",
    DataBarRule(start_type="num", start_value=0, end_type="max", color=TEAL, showValue=True))
pd_.freeze_panes = "A4"

# =================================================================== LISTAS
ls = wb.create_sheet("Listas")
ls.sheet_view.showGridLines = False
ls.sheet_properties.tabColor = GREY

for col, w in (("A", 2.5), ("B", 40), ("C", 2.5), ("D", 32), ("E", 2.5),
               ("F", 16), ("G", 2.5), ("H", 16), ("I", 2.5), ("J", 24), ("K", 2.5)):
    ls.column_dimensions[col].width = w

ls.row_dimensions[1].height = 32
ls.merge_cells("B1:J1")
ls["B1"] = "LISTAS DOS MENUS SUSPENSOS"
ls["B1"].font = f(15, True, "FFFFFF")
ls["B1"].alignment = left(indent=1)
paint(ls, 2, 1, 10, 1, fill(INK))

ls.row_dimensions[2].height = 20
ls.merge_cells("B2:J2")
ls["B2"] = ("Edite, acrescente ou renomeie itens aqui — os menus da aba Gastos e as tabelas do Painel "
            "acompanham automaticamente. Não deixe linhas em branco no meio de uma lista.")
ls["B2"].font = f(9, False, GREY, italic=True)
ls["B2"].alignment = left(indent=1)
paint(ls, 2, 2, 10, 2, fill(BAND))
ls.row_dimensions[3].height = 10

blocos = [
    (2, "CATEGORIAS", CATEGORIAS, 20),
    (4, "FORMAS DE PAGAMENTO", PAGAMENTOS, 20),
    (6, "STATUS", STATUS, 20),
    (8, "QUEM PAGOU", PESSOAS, 20),
    (10, "CIDADES / LOCAIS", CIDADES, 20),
]
for col, titulo, itens, capacidade in blocos:
    cell = ls.cell(row=4, column=col, value=titulo)
    cell.font = f(9, True, "FFFFFF")
    cell.fill = fill("3B4A6B")
    cell.alignment = center()
    for i in range(capacidade):
        rr = 5 + i
        c = ls.cell(row=rr, column=col, value=itens[i] if i < len(itens) else None)
        c.font = f(10, False, INPUT_FG if i < len(itens) else INK)
        c.fill = fill(INPUT_BG if i < len(itens) else "FFFFFF")
        c.alignment = left(indent=1)
        c.border = Border(bottom=thin)
        ls.row_dimensions[rr].height = 18
    outline(ls, col, 4, col, 4 + capacidade)
ls.row_dimensions[4].height = 22

# =================================================================== nomes definidos
def dinamica(col):
    """Intervalo que cresce sozinho conforme itens sao adicionados na aba Listas."""
    return (f"OFFSET(Listas!${col}$5,0,0,"
            f"MAX(1,COUNTA(Listas!${col}$5:${col}$24)),1)")


nomes = {
    "CATEGORIAS": dinamica("B"),
    "PAGAMENTOS": dinamica("D"),
    "STATUS":     dinamica("F"),
    "PESSOAS":    dinamica("H"),
    "CIDADES":    dinamica("J"),
    "MOEDAS":     "Painel!$E$8:$E$11",
    "TAXAS":      "Painel!$F$8:$F$11",
}
for nome, ref in nomes.items():
    wb.defined_names.add(DefinedName(nome, attr_text=ref))

# =================================================================== impressao
from openpyxl.worksheet.properties import PageSetupProperties

for ws, titulos in ((pa, None), (gs, "3:3"), (pd_, "3:3"), (ls, None)):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if titulos:
        ws.print_title_rows = titulos

pa.print_area = f"A1:I{guia_r + len(linhas)}"
gs.print_area = "A1:K203"          # 200 lancamentos por padrao na impressao
pd_.print_area = f"A1:I{tot}"
ls.print_area = "A1:K25"

wb.active = 0
wb.save(OUT)
print("ok ->", OUT)
